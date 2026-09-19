"""A site's EQUIPMENT REGISTRY: systems, equipment, point slots, and the schedule
import that fills them.

WHY THIS FILE EXISTS. Building Intelligence knew points and sites and nothing in
between, so a chiller's ΔT design band was a literal in a formula string and kW/TR
could not be computed at all. Migration 0032 gives equipment a home. What is
asserted here is that the home is honest:

  * every word — system kind, equipment class, slot, design fact — is from a
    CLOSED vocabulary, refused at the edge otherwise;
  * a slot binds to a point by the gateway's TAGS, one point to one slot;
  * design facts are typed: numbers are numbers, a band is min < max, and the
    unit travels with the value;
  * the registry is tenant-scoped and site-confined like every sites table;
  * every write publishes the whole equipment, so a mirror can upsert blindly;
  * an I/O schedule import reports what it would do, does exactly that when
    asked, never half-writes, and skips an ambiguous row instead of choosing.

Every tag, name and number below is a TEST INPUT, not a statement about any real
plant — none of them is the live estate's.
"""

import io

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError

from app.sites.infrastructure.models import EquipmentPointSlot, SiteEquipment, SiteSystem
from app.sites.site.models import Site
from app.tenancy.models import Tenant
from conftest import api_client, bearer, make_role, make_user

pytestmark = pytest.mark.asyncio

PREFIX = "/api/v1"
FIXTURE = __import__("pathlib").Path(__file__).parent / "fixtures" / "io_schedule_example.xlsx"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def _tenant(db, slug: str, *, analytics: bool = True) -> Tenant:
    # The registry is Building Intelligence configuration and rides on the BI
    # module; a tenant that never bought it is refused at the router.
    t = Tenant(name=slug, slug=slug, status="active",
               features={"analytics": analytics}, limits={})
    db.add(t)
    await db.commit()
    await db.refresh(t)
    return t


async def _site(db, tenant, site_id="site-1", *, active=True):
    db.add(Site(site_id=site_id, tenant_id=tenant.id, name=f"Tower {site_id}",
                site_type="building", is_active=active))
    await db.commit()


async def _operator(db, tenant, email, perms=("bi.read", "bi.manage"), site_ids=None):
    role = await make_role(db, f"Role-{email}", list(perms))
    user = await make_user(db, email, role)
    user.tenant_id = tenant.id
    if site_ids is not None:
        user.site_ids = list(site_ids)
    await db.commit()
    await db.refresh(user, attribute_names=["role"])
    return user


def _base(site_id="site-1") -> str:
    return f"{PREFIX}/sites/{site_id}/infrastructure"


async def _system(c, user, name="Loop X", kind="chw_plant", site_id="site-1") -> str:
    r = await c.post(f"{_base(site_id)}/systems", json={"name": name, "kind": kind},
                     headers=bearer(user))
    assert r.status_code == 201, r.text
    return r.json()["system_id"]


async def _chiller(c, user, system_id, tag="TEST-CH-1", site_id="site-1", **extra):
    body = {"system_id": system_id, "tag": tag, "equipment_class": "chiller", **extra}
    r = await c.post(f"{_base(site_id)}/equipment", json=body, headers=bearer(user))
    assert r.status_code == 201, r.text
    return r.json()


async def _count(db, table: str) -> int:
    return (await db.execute(text(f"SELECT count(*) FROM {table}"))).scalar()


def _capture(monkeypatch):
    """Every published event, from both writers of the registry."""
    from app.sites.infrastructure import schedule_import as imp_mod
    from app.sites.infrastructure import service as svc_mod

    seen = []

    async def _emit(tenant_id, entity, event, payload):
        seen.append((entity, event, payload))

    monkeypatch.setattr(svc_mod, "emit", _emit)
    monkeypatch.setattr(imp_mod, "emit", _emit)
    return seen


# ── CRUD, tenant-scoped ──────────────────────────────────────────────────────


async def test_a_chiller_with_facts_and_slots_is_created_and_read_back(app, db):
    t = await _tenant(db, "inf-crud")
    await _site(db, t)
    user = await _operator(db, t, "crud@x.io")

    async with api_client(app) as c:
        sid = await _system(c, user)
        made = await _chiller(
            c, user, sid, name="Test chiller",
            design={"tr": 111, "design_dt_min": 2, "design_dt_max": 3, "make": "TestMake"},
            slots=[{"slot": "chws", "device_tag": "test-dev", "point_tag": "test-leave"},
                   {"slot": "run_status"}],
        )
        tree = await c.get(_base(), headers=bearer(user))
    assert tree.status_code == 200, tree.text
    [system] = tree.json()["systems"]
    assert (system["name"], system["kind"]) == ("Loop X", "chw_plant")
    [eq] = system["equipment"]
    assert eq["equipment_id"] == made["equipment_id"]
    assert eq["design"] == {"tr": 111, "design_dt_min": 2, "design_dt_max": 3, "make": "TestMake"}
    assert eq["design_units"] == {"tr": "TR", "design_dt_min": "K", "design_dt_max": "K"}
    slots = {s["slot"]: s for s in eq["slots"]}
    assert slots["chws"] == {"slot": "chws", "device_tag": "test-dev",
                             "point_tag": "test-leave", "bound": True}
    assert slots["run_status"]["bound"] is False


async def test_equipment_is_renamed_and_moved_within_its_site(app, db):
    t = await _tenant(db, "inf-move")
    await _site(db, t)
    user = await _operator(db, t, "move@x.io")

    async with api_client(app) as c:
        a = await _system(c, user, "Loop X")
        b = await _system(c, user, "Loop Z")
        eq = await _chiller(c, user, a)
        r = await c.patch(f"{_base()}/equipment/{eq['equipment_id']}",
                          json={"tag": "TEST-CH-9", "system_id": b}, headers=bearer(user))
    assert r.status_code == 200, r.text
    assert (r.json()["tag"], r.json()["system_id"]) == ("TEST-CH-9", b)


async def test_equipment_cannot_move_into_a_system_of_the_wrong_kind(app, db):
    t = await _tenant(db, "inf-wrongkind")
    await _site(db, t)
    user = await _operator(db, t, "wk@x.io")

    async with api_client(app) as c:
        plant = await _system(c, user, "Loop X")
        power = await _system(c, user, "Power", kind="power")
        eq = await _chiller(c, user, plant)
        r = await c.patch(f"{_base()}/equipment/{eq['equipment_id']}",
                          json={"system_id": power}, headers=bearer(user))
    assert r.status_code == 422, r.text


async def test_deleting_equipment_takes_its_slots(app, db):
    t = await _tenant(db, "inf-del")
    await _site(db, t)
    user = await _operator(db, t, "del@x.io")

    async with api_client(app) as c:
        sid = await _system(c, user)
        eq = await _chiller(c, user, sid, slots=[{"slot": "kw", "device_tag": "d",
                                                  "point_tag": "p"}])
        r = await c.delete(f"{_base()}/equipment/{eq['equipment_id']}", headers=bearer(user))
    assert r.status_code == 204, r.text
    assert await _count(db, "site_equipment") == 0
    assert await _count(db, "equipment_point_slots") == 0


async def test_deleting_a_system_takes_its_equipment_and_says_so_per_unit(app, db, monkeypatch):
    """A mirror that listens only to equipment must still hear each one go."""
    seen = _capture(monkeypatch)
    t = await _tenant(db, "inf-delsys")
    await _site(db, t)
    user = await _operator(db, t, "delsys@x.io")

    async with api_client(app) as c:
        sid = await _system(c, user)
        e1 = await _chiller(c, user, sid, tag="TEST-CH-1",
                            slots=[{"slot": "kw", "device_tag": "d", "point_tag": "p"}])
        e2 = await _chiller(c, user, sid, tag="TEST-CH-2")
        seen.clear()
        r = await c.delete(f"{_base()}/systems/{sid}", headers=bearer(user))
    assert r.status_code == 204, r.text
    assert await _count(db, "site_systems") == 0
    assert await _count(db, "site_equipment") == 0
    assert await _count(db, "equipment_point_slots") == 0
    assert [(e, ev) for e, ev, _ in seen] == [
        ("equipment", "deleted"), ("equipment", "deleted"), ("site_system", "deleted")
    ]
    assert {p["equipment_id"] for _, _, p in seen[:2]} == {e1["equipment_id"], e2["equipment_id"]}
    assert sorted(seen[2][2]["equipment_ids"]) == sorted([e1["equipment_id"], e2["equipment_id"]])


async def test_another_tenant_cannot_read_or_write_the_registry(app, db):
    mine = await _tenant(db, "inf-mine")
    theirs = await _tenant(db, "inf-theirs")
    await _site(db, theirs, "their-site")
    owner = await _operator(db, theirs, "owner@x.io")
    intruder = await _operator(db, mine, "intruder@x.io")

    async with api_client(app) as c:
        sid = await _system(c, owner, site_id="their-site")
        eq = await _chiller(c, owner, sid, site_id="their-site")
        base = _base("their-site")
        reads = [
            await c.get(base, headers=bearer(intruder)),
            await c.get(f"{base}/equipment/{eq['equipment_id']}", headers=bearer(intruder)),
        ]
        writes = [
            await c.post(f"{base}/systems", json={"name": "Mine", "kind": "power"},
                         headers=bearer(intruder)),
            await c.delete(f"{base}/equipment/{eq['equipment_id']}", headers=bearer(intruder)),
            await c.put(f"{base}/equipment/{eq['equipment_id']}/design",
                        json={"design": {}}, headers=bearer(intruder)),
        ]
    assert [r.status_code for r in reads + writes] == [404] * 5
    assert await _count(db, "site_equipment") == 1


async def test_equipment_is_not_reachable_through_another_sites_url(app, db):
    """The equipment id alone is not the address; the site in the URL must own it."""
    t = await _tenant(db, "inf-cross")
    await _site(db, t, "site-1")
    await _site(db, t, "site-2")
    user = await _operator(db, t, "cross@x.io")

    async with api_client(app) as c:
        sid = await _system(c, user)
        eq = await _chiller(c, user, sid)
        r = await c.get(f"{_base('site-2')}/equipment/{eq['equipment_id']}",
                        headers=bearer(user))
    assert r.status_code == 404, r.text


async def test_a_site_confined_user_sees_only_their_sites_registry(app, db):
    t = await _tenant(db, "inf-confined")
    await _site(db, t, "site-1")
    await _site(db, t, "site-2")
    user = await _operator(db, t, "conf@x.io", site_ids=["site-1"])

    async with api_client(app) as c:
        ok = await c.get(_base("site-1"), headers=bearer(user))
        hidden = await c.get(_base("site-2"), headers=bearer(user))
    assert ok.status_code == 200, ok.text
    assert hidden.status_code == 404, hidden.text


async def test_reading_needs_bi_read_and_writing_needs_bi_manage(app, db):
    t = await _tenant(db, "inf-perm")
    await _site(db, t)
    reader = await _operator(db, t, "reader@x.io", perms=("bi.read",))
    nobody = await _operator(db, t, "nobody@x.io", perms=("devices.read",))

    async with api_client(app) as c:
        read = await c.get(_base(), headers=bearer(reader))
        write = await c.post(f"{_base()}/systems", json={"name": "L", "kind": "power"},
                             headers=bearer(reader))
        blind = await c.get(_base(), headers=bearer(nobody))
    assert (read.status_code, write.status_code, blind.status_code) == (200, 403, 403)


async def test_nothing_is_added_to_a_deleted_site(app, db):
    """Equipment written into a soft-deleted building would reappear, unasked,
    when the building is restored."""
    t = await _tenant(db, "inf-inactive")
    await _site(db, t, active=False)
    user = await _operator(db, t, "inactive@x.io")

    async with api_client(app) as c:
        r = await c.post(f"{_base()}/systems", json={"name": "L", "kind": "power"},
                         headers=bearer(user))
    assert r.status_code == 404, r.text


# ── the closed vocabulary ────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "path, body, says, why",
    [
        ("systems", {"name": "L", "kind": "hvac"}, "unknown system kind",
         "a kind nobody defined"),
        ("equipment", {"tag": "T-1", "equipment_class": "boiler"}, "unknown equipment class",
         "a class nobody defined"),
        ("equipment", {"tag": "T-1", "equipment_class": "Chiller"}, "unknown equipment class",
         "the API matches exactly"),
        ("equipment", {"tag": "T-1", "equipment_class": "dg_set"}, "cannot sit in",
         "a DG set is not part of a chilled-water loop"),
        ("equipment", {"tag": "T-1", "equipment_class": "chiller",
                       "slots": [{"slot": "fuel_level"}]}, "has no 'fuel_level' slot",
         "a chiller has no fuel level"),
        ("equipment", {"tag": "T-1", "equipment_class": "chiller",
                       "slots": [{"slot": "chws_temp"}]}, "unknown slot",
         "a slot nobody defined"),
        ("equipment", {"tag": "T-1", "equipment_class": "chw_primary_pump",
                       "design": {"tr": 10}}, "has no 'tr' design fact", "a pump has no TR"),
        ("equipment", {"tag": "T-1", "equipment_class": "chiller",
                       "design": {"capacity": 10}}, "unknown design fact",
         "a fact nobody defined"),
    ],
)
async def test_a_word_outside_the_vocabulary_is_refused(app, db, path, body, says, why):
    t = await _tenant(db, f"inf-vocab-{abs(hash(why)) % 10**8}")
    await _site(db, t)
    user = await _operator(db, t, f"v{abs(hash(why)) % 10**8}@x.io")

    async with api_client(app) as c:
        sid = await _system(c, user)
        if path == "equipment":
            body = {"system_id": sid, **body}
        r = await c.post(f"{_base()}/{path}", json=body, headers=bearer(user))
    assert r.status_code == 422, why
    # The sentence, too: it is what the designer shows, and it tells a word the
    # platform has never heard of from one the class does not carry.
    assert says in r.json()["error"]["message"], (why, r.json())
    assert await _count(db, "site_equipment") == 0


async def test_a_slot_the_class_does_not_have_cannot_be_bound_later(app, db):
    t = await _tenant(db, "inf-vocab-put")
    await _site(db, t)
    user = await _operator(db, t, "vput@x.io")

    async with api_client(app) as c:
        sid = await _system(c, user)
        eq = await _chiller(c, user, sid)
        r = await c.put(f"{_base()}/equipment/{eq['equipment_id']}/slots/fuel_level",
                        json={"device_tag": "d", "point_tag": "p"}, headers=bearer(user))
    assert r.status_code == 422, r.text
    assert await _count(db, "equipment_point_slots") == 0


async def test_the_class_of_existing_equipment_cannot_be_patched(app, db):
    """Its slots and facts were admitted against the class."""
    t = await _tenant(db, "inf-classfix")
    await _site(db, t)
    user = await _operator(db, t, "cf@x.io")

    async with api_client(app) as c:
        sid = await _system(c, user)
        eq = await _chiller(c, user, sid)
        r = await c.patch(f"{_base()}/equipment/{eq['equipment_id']}",
                          json={"equipment_class": "cooling_tower"}, headers=bearer(user))
    assert r.status_code == 422, r.text


async def test_the_vocabulary_is_published_for_the_designer(app, db):
    t = await _tenant(db, "inf-vocdoc")
    user = await _operator(db, t, "vd@x.io")
    async with api_client(app) as c:
        r = await c.get(f"{PREFIX}/site-infrastructure/vocabulary", headers=bearer(user))
    assert r.status_code == 200, r.text
    doc = r.json()
    chiller = next(e for e in doc["equipment_classes"] if e["key"] == "chiller")
    assert "design_dt_min" in chiller["facts"] and "chws" in chiller["slots"]
    tr = next(f for f in doc["design_facts"] if f["key"] == "tr")
    assert (tr["type"], tr["unit"]) == ("number", "TR")


# ── binding by tag ───────────────────────────────────────────────────────────


async def test_a_slot_is_bound_by_device_and_point_tag(app, db):
    t = await _tenant(db, "inf-bind")
    await _site(db, t)
    user = await _operator(db, t, "bind@x.io")

    async with api_client(app) as c:
        sid = await _system(c, user)
        eq = await _chiller(c, user, sid)
        r = await c.put(f"{_base()}/equipment/{eq['equipment_id']}/slots/chwr",
                        json={"device_tag": "  test dev 01 ", "point_tag": "TEST_Pt Enter"},
                        headers=bearer(user))
    assert r.status_code == 200, r.text
    row = (await db.execute(text(
        "SELECT device_tag, point_tag FROM equipment_point_slots WHERE slot = 'chwr'"
    ))).one()
    # Outer whitespace trimmed; case and inner spaces kept exactly.
    assert tuple(row) == ("test dev 01", "TEST_Pt Enter")


async def test_a_slot_is_unbound_with_two_nulls_and_stays_declared(app, db):
    t = await _tenant(db, "inf-unbind")
    await _site(db, t)
    user = await _operator(db, t, "unbind@x.io")

    async with api_client(app) as c:
        sid = await _system(c, user)
        eq = await _chiller(c, user, sid, slots=[{"slot": "kw", "device_tag": "d",
                                                  "point_tag": "p"}])
        r = await c.put(f"{_base()}/equipment/{eq['equipment_id']}/slots/kw",
                        json={"device_tag": None, "point_tag": None}, headers=bearer(user))
    assert r.status_code == 200, r.text
    assert r.json()["slots"] == [{"slot": "kw", "device_tag": None, "point_tag": None,
                                  "bound": False}]
    nulls = (await db.execute(text(
        "SELECT count(*) FROM equipment_point_slots WHERE device_tag IS NULL "
        "AND point_tag IS NULL"
    ))).scalar()
    assert nulls == 1


async def test_half_a_binding_is_refused(app, db):
    t = await _tenant(db, "inf-half")
    await _site(db, t)
    user = await _operator(db, t, "half@x.io")

    async with api_client(app) as c:
        sid = await _system(c, user)
        eq = await _chiller(c, user, sid)
        r = await c.put(f"{_base()}/equipment/{eq['equipment_id']}/slots/kw",
                        json={"device_tag": "d"}, headers=bearer(user))
    assert r.status_code == 422, r.text


async def test_the_database_refuses_half_a_binding(db):
    """The constraint, not the validator: the importer, psql and a future
    migration are writers too."""
    db.add(EquipmentPointSlot(site_id="s", equipment_id="e", slot="kw", device_tag="d"))
    with pytest.raises(IntegrityError):
        await db.commit()
    await db.rollback()


async def test_one_point_cannot_feed_two_slots(app, db):
    """A kW point on two chillers is plant kW counted twice."""
    t = await _tenant(db, "inf-twice")
    await _site(db, t)
    user = await _operator(db, t, "twice@x.io")

    async with api_client(app) as c:
        sid = await _system(c, user)
        await _chiller(c, user, sid, tag="TEST-CH-1",
                       slots=[{"slot": "kw", "device_tag": "d", "point_tag": "p"}])
        eq2 = await _chiller(c, user, sid, tag="TEST-CH-2")
        r = await c.put(f"{_base()}/equipment/{eq2['equipment_id']}/slots/kw",
                        json={"device_tag": "d", "point_tag": "p"}, headers=bearer(user))
    assert r.status_code == 409, r.text
    assert "TEST-CH-1.kw" in r.json()["error"]["message"]


async def test_the_same_tags_in_another_tenant_are_another_point(app, db):
    """Two tenants' gateways may spell a tag the same way; the binding is per tenant."""
    for slug in ("inf-ta", "inf-tb"):
        t = await _tenant(db, slug)
        await _site(db, t, f"site-{slug}")
        user = await _operator(db, t, f"{slug}@x.io")
        async with api_client(app) as c:
            sid = await _system(c, user, site_id=f"site-{slug}")
            await _chiller(c, user, sid, site_id=f"site-{slug}",
                           slots=[{"slot": "kw", "device_tag": "d", "point_tag": "p"}])
    assert await _count(db, "equipment_point_slots") == 2


# ── design facts ─────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "design, why",
    [
        ({"tr": "111"}, "a number typed as text is not a number"),
        ({"tr": True}, "true is not a capacity"),
        ({"tr": 0}, "a capacity of zero is not a capacity"),
        ({"design_dt_min": 3, "design_dt_max": 3}, "min must be below max"),
        ({"design_dt_min": 4, "design_dt_max": 3}, "an inverted band"),
        ({"design_dt_min": 2}, "half a band is not a band"),
        ({"design_dt_min": -1, "design_dt_max": 3}, "a negative ΔT bound"),
        ({"make": 12}, "make is text"),
    ],
)
async def test_design_facts_are_typed_and_the_band_is_a_band(app, db, design, why):
    t = await _tenant(db, f"inf-df-{abs(hash(why)) % 10**8}")
    await _site(db, t)
    user = await _operator(db, t, f"df{abs(hash(why)) % 10**8}@x.io")

    async with api_client(app) as c:
        sid = await _system(c, user)
        created = await c.post(f"{_base()}/equipment",
                               json={"system_id": sid, "tag": "T-1",
                                     "equipment_class": "chiller", "design": design},
                               headers=bearer(user))
        eq = await _chiller(c, user, sid, tag="T-2")
        put = await c.put(f"{_base()}/equipment/{eq['equipment_id']}/design",
                          json={"design": design}, headers=bearer(user))
    assert created.status_code == 422, why
    assert put.status_code == 422, why


async def test_the_design_put_is_a_set_and_null_clears(app, db):
    t = await _tenant(db, "inf-dset")
    await _site(db, t)
    user = await _operator(db, t, "dset@x.io")

    async with api_client(app) as c:
        sid = await _system(c, user)
        eq = await _chiller(c, user, sid, design={"tr": 111, "make": "TestMake"})
        r = await c.put(f"{_base()}/equipment/{eq['equipment_id']}/design",
                        json={"design": {"tr": None, "design_dt_min": 2.5,
                                         "design_dt_max": 3.5}},
                        headers=bearer(user))
    assert r.status_code == 200, r.text
    assert r.json()["design"] == {"design_dt_min": 2.5, "design_dt_max": 3.5}
    assert r.json()["design_units"] == {"design_dt_min": "K", "design_dt_max": "K"}


async def test_equipment_with_no_facts_stores_an_empty_object(db):
    """`{}` is the one spelling of "nothing recorded", and the DATABASE supplies
    it — a writer that says nothing about design gets the same row the API makes."""
    await db.execute(text(
        "INSERT INTO site_systems (system_id, site_id, name, kind, created_at, updated_at) "
        "VALUES ('s', 'x', 'L', 'chw_plant', '2026-01-01', '2026-01-01')"
    ))
    await db.execute(text(
        "INSERT INTO site_equipment (equipment_id, site_id, system_id, tag, "
        "equipment_class, created_at, updated_at) "
        "VALUES ('e', 'x', 's', 'T', 'chiller', '2026-01-01', '2026-01-01')"
    ))
    await db.commit()
    design = (await db.execute(text("SELECT design FROM site_equipment"))).scalar()
    assert design == "{}"


async def test_a_python_none_design_is_never_stored_as_json_null(db):
    """The 0031 trap on a NOT NULL column: without `none_as_null`, None becomes
    the JSON scalar `null` — a value — and NOT NULL waves it through. With it, a
    None on insert is "no value" and the column's own `{}` is written; on update
    it is SQL NULL and refused."""
    db.add(SiteSystem(system_id="s", site_id="x", name="L", kind="chw_plant"))
    db.add(SiteEquipment(equipment_id="e", site_id="x", system_id="s", tag="T",
                         equipment_class="chiller", design=None))
    await db.commit()
    assert (await db.execute(text("SELECT design FROM site_equipment"))).scalar() == "{}"

    row = await db.get(SiteEquipment, "e")
    row.design = None
    with pytest.raises(IntegrityError):
        await db.commit()
    await db.rollback()


# ── what is published ────────────────────────────────────────────────────────


async def test_every_write_publishes_one_event_carrying_the_whole_equipment(
    app, db, monkeypatch
):
    seen = _capture(monkeypatch)
    t = await _tenant(db, "inf-events")
    await _site(db, t)
    user = await _operator(db, t, "events@x.io")

    async with api_client(app) as c:
        sid = await _system(c, user)
        eq = await _chiller(c, user, sid, design={"tr": 111})
        base = f"{_base()}/equipment/{eq['equipment_id']}"
        await c.patch(base, json={"name": "Renamed"}, headers=bearer(user))
        await c.put(f"{base}/design", json={"design": {"tr": 111, "design_dt_min": 2,
                                                       "design_dt_max": 3}},
                    headers=bearer(user))
        await c.put(f"{base}/slots/chws", json={"device_tag": "d", "point_tag": "p"},
                    headers=bearer(user))
        await c.delete(f"{base}/slots/chws", headers=bearer(user))
        await c.patch(f"{_base()}/systems/{sid}", json={"name": "Loop Q"},
                      headers=bearer(user))
        await c.delete(base, headers=bearer(user))

    assert [(e, ev) for e, ev, _ in seen] == [
        ("site_system", "created"),
        ("equipment", "created"),
        ("equipment", "updated"),
        ("equipment", "design_updated"),
        ("equipment", "slot_set"),
        ("equipment", "slot_removed"),
        ("site_system", "updated"),
        ("equipment", "deleted"),
    ]
    design_event = seen[3][2]
    assert design_event["design"] == {"tr": 111, "design_dt_min": 2, "design_dt_max": 3}
    assert design_event["design_units"] == {"tr": "TR", "design_dt_min": "K",
                                            "design_dt_max": "K"}
    assert design_event["system_kind"] == "chw_plant"
    assert design_event["source"] == "designer"
    # The whole equipment, read back after commit: the slot event names the slot
    # AND carries the binding.
    assert seen[4][2]["slots"] == [{"slot": "chws", "device_tag": "d", "point_tag": "p"}]
    assert seen[5][2]["slots"] == []
    assert seen[7][2] == {"site_id": "site-1", "system_id": sid,
                          "equipment_id": eq["equipment_id"], "tag": "TEST-CH-1"}


# ── the I/O schedule import ──────────────────────────────────────────────────

HEADER = ["System", "System Kind", "Equipment Tag", "Equipment Class", "Equipment Name",
          "Slot", "Device Tag", "Point Tag", "TR", "Design dT Min (K)", "Design dT Max (K)"]


def _workbook(rows, header=HEADER, sheet="Equipment_Schedule") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(header)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _row(tag="TEST-CH-1", slot=None, device=None, point=None, tr=None, lo=None, hi=None,
         cls="chiller", system="Loop X", kind="chw_plant", name=None):
    return [system, kind, tag, cls, name, slot, device, point, tr, lo, hi]


async def _import(c, user, content: bytes, *, dry_run=None, site_id="site-1"):
    params = {} if dry_run is None else {"dry_run": str(dry_run).lower()}
    return await c.post(f"{_base(site_id)}/import", params=params,
                        files={"file": ("schedule.xlsx", content, XLSX)},
                        headers=bearer(user))


def _without_ids(report: dict) -> dict:
    out = {k: v for k, v in report.items() if k != "dry_run"}
    out["systems"] = [{k: v for k, v in s.items() if k != "system_id"} for s in out["systems"]]
    out["equipment"] = [{k: v for k, v in e.items() if k != "equipment_id"}
                        for e in out["equipment"]]
    return out


async def test_the_fixture_shows_the_layout_and_imports_as_documented(app, db):
    """The committed fixture IS the documented layout: its header row is the
    column list, and it carries one clean plant, one ambiguous chiller, one
    word outside the vocabulary and an extra column that is ignored."""
    header = [c.value for c in next(load_workbook(FIXTURE)["Equipment_Schedule"].iter_rows())]
    assert header[:4] == ["System", "System Kind", "Equipment Tag", "Equipment Class"]

    t = await _tenant(db, "inf-fixture")
    await _site(db, t)
    user = await _operator(db, t, "fx@x.io")
    async with api_client(app) as c:
        r = await _import(c, user, FIXTURE.read_bytes(), dry_run=False)
    assert r.status_code == 200, r.text
    rep = r.json()
    assert rep["ignored_columns"] == ["Remarks"]
    assert [e["tag"] for e in rep["equipment"]] == ["TEST-CH-1", "TEST-CT-1", "TEST-AHU-1"]
    ch = rep["equipment"][0]
    assert ch["design"] == {"make": "TestMake", "model": "TM-1", "tr": 111,
                            "design_dt_min": 2, "design_dt_max": 3}
    assert [s["slot"] for s in ch["slots"]] == ["chws", "chwr", "kw", "run_status"]
    assert rep["equipment"][1]["equipment_class"] == "cooling_tower"  # "Cooling Tower"
    reasons = {(s["equipment_tag"], s["reason"]) for s in rep["skipped"]}
    assert reasons == {("TEST-CH-2", "ambiguous"), ("TEST-BLR-1", "invalid")}
    assert await _count(db, "site_equipment") == 3


async def test_a_dry_run_writes_nothing_and_the_real_run_does_exactly_what_it_showed(
    app, db, monkeypatch
):
    seen = _capture(monkeypatch)
    t = await _tenant(db, "inf-dry")
    await _site(db, t)
    user = await _operator(db, t, "dry@x.io")
    content = FIXTURE.read_bytes()

    async with api_client(app) as c:
        dry = await _import(c, user, content, dry_run=True)
        counts_after_dry = [await _count(db, n) for n in
                            ("site_systems", "site_equipment", "equipment_point_slots")]
        events_after_dry = len(seen)
        real = await _import(c, user, content, dry_run=False)

    assert dry.status_code == 200 and real.status_code == 200, (dry.text, real.text)
    assert counts_after_dry == [0, 0, 0]
    assert events_after_dry == 0
    assert dry.json()["dry_run"] is True and real.json()["dry_run"] is False
    assert _without_ids(dry.json()) == _without_ids(real.json())
    assert all(e["equipment_id"] for e in real.json()["equipment"])
    assert [await _count(db, n) for n in
            ("site_systems", "site_equipment", "equipment_point_slots")] == [2, 3, 6]


async def test_an_import_without_the_flag_is_a_dry_run(app, db):
    t = await _tenant(db, "inf-default")
    await _site(db, t)
    user = await _operator(db, t, "default@x.io")
    async with api_client(app) as c:
        r = await _import(c, user, _workbook([_row()]))
    assert r.status_code == 200, r.text
    assert r.json()["dry_run"] is True
    assert await _count(db, "site_equipment") == 0


async def test_a_real_import_publishes_one_event_per_row_it_made(app, db, monkeypatch):
    seen = _capture(monkeypatch)
    t = await _tenant(db, "inf-impev")
    await _site(db, t)
    user = await _operator(db, t, "impev@x.io")
    async with api_client(app) as c:
        r = await _import(c, user, FIXTURE.read_bytes(), dry_run=False)
    assert r.status_code == 200, r.text
    assert [(e, ev) for e, ev, _ in seen] == [
        ("site_system", "created"), ("site_system", "created"),
        ("equipment", "created"), ("equipment", "created"), ("equipment", "created"),
    ]
    ch = seen[2][2]
    assert ch["source"] == "schedule_import"
    assert ch["design_units"]["tr"] == "TR"
    assert len(ch["slots"]) == 4


@pytest.mark.parametrize(
    "rows, skipped_tags, why",
    [
        ([_row(tr=100), _row(slot="kw", device="d", point="p", tr=200)],
         {"TEST-CH-1"}, "one chiller, two capacities"),
        ([_row(cls="chiller"), _row(cls="cooling_tower", slot="speed")],
         {"TEST-CH-1"}, "one tag, two classes"),
        ([_row(slot="kw", device="d", point="p1"), _row(slot="kw", device="d", point="p2")],
         {"TEST-CH-1"}, "one slot, bound two ways"),
        ([_row("TEST-CH-1", slot="kw", device="d", point="p"),
          _row("TEST-CH-2", slot="kw", device="d", point="p")],
         {"TEST-CH-1", "TEST-CH-2"}, "one point, two slots"),
        ([_row("TEST-CH-1"), _row("TEST-AHU-1", cls="ahu", kind="air_handling")],
         {"TEST-CH-1", "TEST-AHU-1"}, "one system name, two kinds"),
    ],
)
async def test_an_ambiguous_row_is_reported_and_skipped_never_resolved(
    app, db, rows, skipped_tags, why
):
    t = await _tenant(db, f"inf-amb-{abs(hash(why)) % 10**8}")
    await _site(db, t)
    user = await _operator(db, t, f"amb{abs(hash(why)) % 10**8}@x.io")
    async with api_client(app) as c:
        r = await _import(c, user, _workbook(rows), dry_run=False)
    assert r.status_code == 200, r.text
    rep = r.json()
    ambiguous = [s for s in rep["skipped"] if s["reason"] == "ambiguous"]
    assert {s["equipment_tag"] for s in ambiguous} == skipped_tags, why
    # Every row involved is skipped — keeping either one is the guess.
    assert {s["row"] for s in ambiguous} == set(range(2, 2 + len(rows))), why
    # Nothing the ambiguity touched was written.
    bound = (await db.execute(text(
        "SELECT count(*) FROM equipment_point_slots WHERE device_tag IS NOT NULL"
    ))).scalar()
    assert bound == 0, why


@pytest.mark.parametrize(
    "row, why",
    [
        (_row(cls="boiler"), "a class nobody defined"),
        (_row(slot="fuel_level"), "a slot the class does not have"),
        (_row(slot="kw", device="d"), "half a binding"),
        (_row(tr="111"), "a number typed as text"),
        (_row(lo=3, hi=2), "an inverted band"),
        (_row(system=None), "a required cell left blank"),
    ],
)
async def test_an_invalid_row_is_reported_with_its_row_number(app, db, row, why):
    t = await _tenant(db, f"inf-inv-{abs(hash(why)) % 10**8}")
    await _site(db, t)
    user = await _operator(db, t, f"inv{abs(hash(why)) % 10**8}@x.io")
    good = _row("TEST-CT-1", cls="cooling_tower", slot="speed")
    bad = row[:2] + ["TEST-BAD-1"] + row[3:]
    async with api_client(app) as c:
        r = await _import(c, user, _workbook([good, bad]), dry_run=False)
    assert r.status_code == 200, r.text
    rep = r.json()
    assert [(s["row"], s["reason"]) for s in rep["skipped"]] == [(3, "invalid")], why
    # The good row is not held hostage by the bad one.
    assert [e["tag"] for e in rep["equipment"]] == ["TEST-CT-1"], why


async def test_a_bad_fact_skips_its_row_not_the_equipments_other_rows(app, db):
    """The row is what is refused. The same chiller's other rows said nothing
    wrong, and the chiller is created from them — without the unreadable fact,
    which is "not recorded", not a guess at what "111" meant."""
    t = await _tenant(db, "inf-rowfact")
    await _site(db, t)
    user = await _operator(db, t, "rowfact@x.io")
    rows = [_row(slot="kw", device="d", point="p"), _row(tr="111")]
    async with api_client(app) as c:
        r = await _import(c, user, _workbook(rows), dry_run=False)
    assert r.status_code == 200, r.text
    rep = r.json()
    assert [(s["row"], s["reason"]) for s in rep["skipped"]] == [(3, "invalid")]
    [eq] = rep["equipment"]
    assert eq["design"] == {}
    assert eq["slots"] == [{"slot": "kw", "device_tag": "d", "point_tag": "p"}]


async def test_equipment_already_in_the_registry_is_not_imported_over(app, db):
    t = await _tenant(db, "inf-exists")
    await _site(db, t)
    user = await _operator(db, t, "exists@x.io")
    content = _workbook([_row(tr=100)])
    async with api_client(app) as c:
        first = await _import(c, user, content, dry_run=False)
        second = await _import(c, user, content, dry_run=False)
    assert first.status_code == second.status_code == 200
    assert [s["reason"] for s in second.json()["skipped"]] == ["exists"]
    assert second.json()["counts"]["equipment_created"] == 0
    assert await _count(db, "site_equipment") == 1


async def test_a_point_already_bound_in_the_registry_is_a_conflict(app, db):
    t = await _tenant(db, "inf-impconf")
    await _site(db, t)
    user = await _operator(db, t, "impconf@x.io")
    async with api_client(app) as c:
        sid = await _system(c, user)
        await _chiller(c, user, sid, tag="TEST-CH-0",
                       slots=[{"slot": "kw", "device_tag": "d", "point_tag": "p"}])
        r = await _import(c, user, _workbook([_row(slot="kw", device="d", point="p")]),
                          dry_run=False)
    assert r.status_code == 200, r.text
    assert [s["reason"] for s in r.json()["skipped"]] == ["conflict"]
    assert r.json()["equipment"][0]["slots"] == []


async def test_an_existing_system_of_the_same_kind_is_reused(app, db):
    t = await _tenant(db, "inf-reuse")
    await _site(db, t)
    user = await _operator(db, t, "reuse@x.io")
    async with api_client(app) as c:
        sid = await _system(c, user, "Loop X")
        r = await _import(c, user, _workbook([_row()]), dry_run=False)
    assert r.status_code == 200, r.text
    assert r.json()["systems"] == [{"name": "Loop X", "kind": "chw_plant", "reused": True,
                                    "system_id": sid}]
    assert await _count(db, "site_systems") == 1


async def test_a_real_import_that_fails_to_commit_writes_nothing(app, db):
    """Never half-written: the plan is applied in ONE commit. A point bound by
    someone else between plan and commit fails the whole import, and the systems
    and equipment it would have created are not left behind without their slots."""
    from app.core.errors import ConflictError
    from app.sites.infrastructure import schedule_import
    from app.sites.infrastructure.service import InfrastructureService
    from app.tenancy.scope import Scope

    t = await _tenant(db, "inf-atomic")
    await _site(db, t)
    svc = InfrastructureService(db, Scope(tenant_id=t.id, is_superadmin=False))
    site = await svc.site("site-1")
    content = _workbook([_row(slot="kw", device="d", point="p"), _row("TEST-CH-2")])
    report = await schedule_import.plan(svc, site, content)
    assert report["counts"]["slots_created"] == 1

    # The race: the same point, bound elsewhere, after the plan was made.
    db.add(SiteSystem(system_id="other", tenant_id=t.id, site_id="site-1", name="Other",
                      kind="power"))
    db.add(SiteEquipment(equipment_id="em", tenant_id=t.id, site_id="site-1",
                         system_id="other", tag="TEST-EM-1", equipment_class="energy_meter"))
    db.add(EquipmentPointSlot(tenant_id=t.id, site_id="site-1", equipment_id="em",
                              slot="kw", device_tag="d", point_tag="p"))
    await db.commit()

    with pytest.raises(ConflictError):
        await schedule_import.apply(svc, site, report, actor=None)
    assert await _count(db, "site_systems") == 1
    assert await _count(db, "site_equipment") == 1


@pytest.mark.parametrize(
    "content, says, why",
    [
        (b"not a workbook", "not an .xlsx workbook", "not a zip at all"),
        (_workbook([_row()], sheet="Sheet1"), "no 'Equipment_Schedule' sheet",
         "no Equipment_Schedule sheet"),
        (_workbook([_row()[:3]], header=HEADER[:3]), "equipment class",
         "no Equipment Class column"),
    ],
    ids=["not-a-zip", "no-sheet", "no-class-column"],
)
async def test_a_file_that_is_not_the_layout_is_refused_whole(app, db, content, says, why):
    t = await _tenant(db, f"inf-file-{abs(hash(why)) % 10**8}")
    await _site(db, t)
    user = await _operator(db, t, f"file{abs(hash(why)) % 10**8}@x.io")
    async with api_client(app) as c:
        r = await _import(c, user, content, dry_run=False)
    assert r.status_code == 422, why
    assert says in r.json()["error"]["message"], (why, r.json())


# ── offboarding ──────────────────────────────────────────────────────────────


async def test_a_tenant_erasure_takes_its_registry_and_only_its_registry(db):
    from app.tenancy.erasure import erase_tenant_data

    doomed = await _tenant(db, "inf-doomed")
    neighbour = await _tenant(db, "inf-neighbour")
    for t, sfx in ((doomed, "d"), (neighbour, "n")):
        db.add(SiteSystem(system_id=f"s{sfx}", tenant_id=t.id, site_id="x", name=f"L{sfx}",
                          kind="chw_plant"))
        db.add(SiteEquipment(equipment_id=f"e{sfx}", tenant_id=t.id, site_id="x",
                             system_id=f"s{sfx}", tag=f"T{sfx}", equipment_class="chiller"))
        db.add(EquipmentPointSlot(tenant_id=t.id, site_id="x", equipment_id=f"e{sfx}",
                                  slot="kw", device_tag="d", point_tag="p"))
    await db.commit()

    await erase_tenant_data(db, doomed.id)
    await db.commit()
    for table in ("site_systems", "site_equipment", "equipment_point_slots"):
        left = (await db.execute(text(f"SELECT tenant_id FROM {table}"))).scalars().all()
        assert [str(v).replace("-", "") for v in left] == [neighbour.id.hex], table



# ── Building Intelligence configuration, not site administration ─────────────


async def test_a_vms_admin_with_sites_update_cannot_describe_a_chiller(app, db):
    """`sites.update` manages buildings for the VMS side. Describing a building's
    plant is BI configuration, and holding the first grants nothing here."""
    t = await _tenant(db, "inf-vms-admin")
    await _site(db, t)
    admin = await _operator(db, t, "vmsadmin@x.io", perms=("sites.read", "sites.update"))

    async with api_client(app) as c:
        read = await c.get(_base(), headers=bearer(admin))
        write = await c.post(f"{_base()}/systems", json={"name": "L", "kind": "power"},
                             headers=bearer(admin))
    assert (read.status_code, write.status_code) == (403, 403)


async def test_a_tenant_without_the_bi_module_is_refused_even_with_bi_manage(app, db):
    t = await _tenant(db, "inf-no-module", analytics=False)
    await _site(db, t)
    user = await _operator(db, t, "nomodule@x.io")

    async with api_client(app) as c:
        tree = await c.get(_base(), headers=bearer(user))
        vocab_r = await c.get(f"{PREFIX}/site-infrastructure/vocabulary", headers=bearer(user))
    assert tree.status_code == 403
    assert tree.json()["error"]["code"] == "FEATURE_DISABLED"
    assert vocab_r.status_code == 403
