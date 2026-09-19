"""I/O schedule import — an ``.xlsx`` into systems, equipment and point slots.

THE LAYOUT IS OURS, AND IT HAS TO BE
------------------------------------
A consultant's I/O schedule is written for the panel builder, not for us. The two
real ones on file (a panel-wise schedule and a "commercial" template) name points
in free text ("CHW Supply Temp", "Run Status"), carry no gateway tags because no
gateway existed when they were drawn, state no TR and no ΔT band, and do not say
which chillers form which loop. Mapping "CHW Supply Temp" to ``chws`` is a guess
this importer will not make on an operator's behalf. So it reads ONE sheet in a
layout defined here, into which a consultant's schedule is transcribed — keeping
their equipment tags — and every word in it is from ``vocabulary.py``:

    Sheet  ``Equipment_Schedule`` — row 1 is the header, one row per slot.

    Column              Required  Meaning
    ------------------  --------  ---------------------------------------------
    System              yes       system name on the site ("Plant A")
    System Kind         yes       vocabulary.SYSTEM_KINDS      (chw_plant …)
    Equipment Tag       yes       the tag on the drawings      (CH-01)
    Equipment Class     yes       vocabulary.EQUIPMENT_CLASSES (chiller …)
    Equipment Name      no        free text
    Slot                no        vocabulary.SLOTS (chws …); blank = the row
                                  only states the equipment and its facts
    Device Tag          no        gateway device tag  } both or neither; blank
    Point Tag           no        gateway point tag   } = slot declared, unbound
    Make, Model         no        text
    TR                  no        number, refrigeration tons
    kW Rated            no        number, kW
    kVA Rated           no        number, kVA
    Design dT Min (K)   no        number, kelvin (a temperature DIFFERENCE)
    Design dT Max (K)   no        number, kelvin

Headers are matched case-insensitively with surrounding space ignored, in any
order; any other column is ignored and reported in ``ignored_columns``, so a
schedule that keeps its Signal Type or Remarks columns still imports. Vocabulary
cells are matched case-insensitively with spaces and hyphens read as underscores
("Cooling Tower" is ``cooling_tower``) — that is spelling, not interpretation.
Numeric facts must be NUMERIC cells: "150" typed as text is refused, not parsed.

WHAT IS REFUSED, AND HOW
------------------------
A row that cannot be read is reported with its sheet row number and skipped;
nothing is guessed to rescue it. Specifically:

  * ``invalid``   — a required cell blank, a word not in the vocabulary, a slot
    the class does not have, half a binding, a fact of the wrong type or range;
  * ``ambiguous`` — two rows disagree: one system given two kinds, one tag given
    two classes / systems / names / values of a fact, one slot bound two ways,
    one point bound to two slots. EVERY row involved is skipped, because
    picking one is exactly the guess;
  * ``exists``    — the equipment tag is already on this site. Import CREATES;
    editing a registry entry is the designer's job, where a human sees it;
  * ``conflict``  — the point is already bound to a slot in the registry;
  * ``duplicate`` — a row repeating an earlier one exactly. Skipped, harmless.

A system that exists on the site under the same name AND kind is reused — the
name is the site-unique key, so that is a match, not a guess. One with the same
name and a different kind is ``ambiguous``.

DRY RUN AND REAL RUN ARE ONE CODE PATH
--------------------------------------
``plan()`` reads the workbook and the registry and writes nothing. A dry run
returns its report; a real run applies the same plan in ONE transaction — every
row it lists, or, if the commit fails, none — and returns the same report with
the new ids filled in. A real run can therefore never do something its dry run
did not show, short of the registry changing in between.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field
from typing import Any

from sqlalchemy import select

from ...core.audit import record as audit_record
from ...core.errors import ValidationError
from ..events import emit
from . import vocabulary as vocab
from .models import EquipmentPointSlot, SiteEquipment, SiteSystem
from .schemas import POINT_TAG_MAX, TAG_MAX

SHEET = "Equipment_Schedule"

#: A workbook is a zip; a cap on the upload bounds what openpyxl is handed.
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
#: Rows past this are refused as a whole rather than silently truncated.
MAX_ROWS = 20_000

# header (normalised) → field
COLUMNS: dict[str, str] = {
    "system": "system",
    "system kind": "kind",
    "equipment tag": "tag",
    "equipment class": "equipment_class",
    "equipment name": "name",
    "slot": "slot",
    "device tag": "device_tag",
    "point tag": "point_tag",
    "make": "make",
    "model": "model",
    "tr": "tr",
    "kw rated": "kw_rated",
    "kva rated": "kva_rated",
    "design dt min (k)": "design_dt_min",
    "design dt max (k)": "design_dt_max",
}
REQUIRED = ("system", "kind", "tag", "equipment_class")
FACT_FIELDS = tuple(k for k in vocab.DESIGN_FACTS)
# The equipment-level attributes rows of one tag must agree on.
_AGREE = ("system", "kind", "equipment_class", "name", *FACT_FIELDS)


def _norm_header(v: Any) -> str:
    return " ".join(str(v).split()).lower() if v is not None else ""


def _word(v: str) -> str:
    return v.strip().lower().replace("-", "_").replace(" ", "_")


@dataclass
class _Row:
    row: int
    system: str
    kind: str
    tag: str
    equipment_class: str
    name: str | None
    slot: str | None
    device_tag: str | None
    point_tag: str | None
    facts: dict[str, Any]

    def agree_value(self, key: str):
        if key in FACT_FIELDS:
            return self.facts.get(key)
        return getattr(self, key)


@dataclass
class _Skip:
    row: int
    equipment_tag: str | None
    slot: str | None
    reason: str
    message: str

    def as_dict(self) -> dict:
        return {"row": self.row, "equipment_tag": self.equipment_tag, "slot": self.slot,
                "reason": self.reason, "message": self.message}


@dataclass
class _PlannedEquipment:
    tag: str
    system: str
    kind: str
    equipment_class: str
    name: str | None
    design: dict[str, Any]
    slots: list[dict[str, Any]] = field(default_factory=list)
    rows: list[int] = field(default_factory=list)


# ── reading the workbook ─────────────────────────────────────────────────────


def _cell_text(v: Any) -> str | None:
    """A text cell, trimmed; blank is None. A NUMBER where text was meant is
    returned as None-with-error by the caller, never stringified."""
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    raise TypeError(type(v).__name__)


def read_rows(content: bytes) -> tuple[list[tuple[int, dict[str, Any]]], list[str]]:
    """(row number, {field: raw cell}) for every non-blank data row, and the
    header cells that were ignored. Refuses the WHOLE file when it is not a
    workbook, lacks the sheet, or lacks a required column — there is no row to
    report against in any of those cases."""
    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise ValidationError("this is not an .xlsx workbook")
    try:
        from openpyxl import load_workbook  # lazily, like the floor-plan converters
    except ImportError as exc:  # pragma: no cover - the image installs [reports]
        raise ValidationError("the spreadsheet reader is not installed on this server") from exc
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any parse failure is the file's fault
        raise ValidationError("this .xlsx could not be read") from exc
    try:
        if SHEET not in wb.sheetnames:
            raise ValidationError(
                f"no {SHEET!r} sheet; this workbook has: {', '.join(wb.sheetnames)}"
            )
        rows = wb[SHEET].iter_rows(values_only=True)
        header = next(rows, None) or ()
        index: dict[str, int] = {}
        ignored: list[str] = []
        for i, h in enumerate(header):
            key = _norm_header(h)
            if not key:
                continue
            fld = COLUMNS.get(key)
            if fld is None:
                ignored.append(str(h).strip())
                continue
            if fld in index:
                raise ValidationError(f"column {str(h).strip()!r} appears twice")
            index[fld] = i
        missing = [
            next(h for h, f in COLUMNS.items() if f == r) for r in REQUIRED if r not in index
        ]
        if missing:
            raise ValidationError("missing required column(s): " + ", ".join(missing))

        out: list[tuple[int, dict[str, Any]]] = []
        for n, values in enumerate(rows, start=2):
            cells = {f: (values[i] if i < len(values) else None) for f, i in index.items()}
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in cells.values()):
                continue
            if len(out) >= MAX_ROWS:
                raise ValidationError(f"more than {MAX_ROWS} rows; split the schedule")
            out.append((n, cells))
        return out, ignored
    finally:
        wb.close()


def _parse(n: int, cells: dict[str, Any]) -> _Row:
    """One row, or ValueError with the sentence the report will carry."""
    text: dict[str, str | None] = {}
    for f in ("system", "kind", "tag", "equipment_class", "name", "slot", "device_tag",
              "point_tag", "make", "model"):
        try:
            text[f] = _cell_text(cells.get(f))
        except TypeError:
            raise ValueError(f"{f.replace('_', ' ')} must be text")
    for f in REQUIRED:
        if text[f] is None:
            raise ValueError(f"{f.replace('_', ' ')} is blank")

    kind = vocab.check_system_kind(_word(text["kind"]))
    cls = vocab.check_equipment_class(_word(text["equipment_class"]))
    vocab.check_class_in_kind(cls, kind)
    if len(text["tag"]) > TAG_MAX:
        raise ValueError(f"equipment tag must be {TAG_MAX} characters or fewer")
    if len(text["system"]) > 100:
        raise ValueError("system name must be 100 characters or fewer")

    slot = vocab.check_slot(cls, _word(text["slot"])) if text["slot"] else None
    dt, pt = text["device_tag"], text["point_tag"]
    if (dt is None) != (pt is None):
        raise ValueError("a binding needs both device tag and point tag, or neither")
    if dt is not None and slot is None:
        raise ValueError("a point is bound to a slot; this row names none")
    if dt is not None and (len(dt) > POINT_TAG_MAX or len(pt) > POINT_TAG_MAX):
        raise ValueError(f"tags must be {POINT_TAG_MAX} characters or fewer")

    facts: dict[str, Any] = {}
    for f in FACT_FIELDS:
        v = text.get(f) if vocab.DESIGN_FACTS[f]["type"] == "text" else cells.get(f)
        if isinstance(v, str) and not v.strip():
            v = None
        if v is not None:
            facts[f] = v
    # Each fact alone here, so the report names the row that carries a bad one;
    # the band (min < max) is checked once per equipment, because its two halves
    # may arrive on different rows.
    for f, v in facts.items():
        facts[f] = vocab.check_fact(cls, f, v)

    return _Row(n, text["system"], kind, text["tag"], cls, text["name"], slot, dt, pt, facts)


# ── planning ────────────────────────────────────────────────────────────────


async def plan(svc, site, content: bytes) -> dict[str, Any]:
    """Everything an import of `content` into `site` would create. Writes nothing."""
    raw, ignored = read_rows(content)
    skipped: list[_Skip] = []
    rows: list[_Row] = []
    for n, cells in raw:
        try:
            rows.append(_parse(n, cells))
        except ValueError as exc:
            tag = cells.get("tag") if isinstance(cells.get("tag"), str) else None
            skipped.append(_Skip(n, tag.strip() if tag else None, None, "invalid", str(exc)))

    def drop(bad: list[_Row], reason: str, message: str) -> None:
        for r in bad:
            skipped.append(_Skip(r.row, r.tag, r.slot, reason, message))
        ids = {id(r) for r in bad}
        rows[:] = [r for r in rows if id(r) not in ids]

    # One system name, one kind — in the sheet, and against the registry.
    existing_systems = {
        s.name: s
        for s in (
            await svc.db.execute(select(SiteSystem).where(SiteSystem.site_id == site.site_id))
        ).scalars().all()
    }
    for name in dict.fromkeys(r.system for r in rows):
        mine = [r for r in rows if r.system == name]
        kinds = sorted({r.kind for r in mine})
        if len(kinds) > 1:
            drop(mine, "ambiguous", f"system {name!r} is given kinds {', '.join(kinds)}")
        elif name in existing_systems and existing_systems[name].kind != kinds[0]:
            drop(mine, "ambiguous",
                 f"system {name!r} already exists here as {existing_systems[name].kind}, "
                 f"not {kinds[0]}")

    # One tag, one piece of equipment.
    existing_tags = set(
        (
            await svc.db.execute(
                select(SiteEquipment.tag).where(SiteEquipment.site_id == site.site_id)
            )
        ).scalars().all()
    )
    planned: dict[str, _PlannedEquipment] = {}
    for tag in dict.fromkeys(r.tag for r in rows):
        mine = [r for r in rows if r.tag == tag]
        if tag in existing_tags:
            drop(mine, "exists", f"{tag} is already in this site's registry")
            continue
        agreed: dict[str, Any] = {}
        clash = None
        for key in _AGREE:
            values = {repr(v): v for v in (r.agree_value(key) for r in mine) if v is not None}
            if len(values) > 1:
                clash = f"{tag} rows disagree on {key.replace('_', ' ')}: " + ", ".join(
                    sorted(values)
                )
                break
            agreed[key] = next(iter(values.values()), None)
        if clash:
            drop(mine, "ambiguous", clash)
            continue
        try:
            design = vocab.check_design(
                agreed["equipment_class"], {f: agreed[f] for f in FACT_FIELDS}
            )
        except ValueError as exc:
            drop(mine, "invalid", str(exc))
            continue
        planned[tag] = _PlannedEquipment(
            tag=tag, system=agreed["system"], kind=agreed["kind"],
            equipment_class=agreed["equipment_class"],
            name=agreed["name"], design=design, rows=[r.row for r in mine],
        )

    # One slot per equipment, one binding per slot, one slot per point.
    slot_rows = [r for r in rows if r.slot is not None and r.tag in planned]
    by_slot: dict[tuple[str, str], list[_Row]] = {}
    for r in slot_rows:
        by_slot.setdefault((r.tag, r.slot), []).append(r)
    kept: list[_Row] = []
    for (tag, slot), mine in by_slot.items():
        bindings = {(r.device_tag, r.point_tag) for r in mine}
        if len(bindings) > 1:
            drop(mine, "ambiguous", f"{tag}.{slot} is bound more than one way")
            continue
        kept.append(mine[0])
        for r in mine[1:]:
            skipped.append(_Skip(r.row, tag, slot, "duplicate", f"repeats row {mine[0].row}"))

    by_point: dict[tuple[str, str], list[_Row]] = {}
    for r in kept:
        if r.device_tag is not None:
            by_point.setdefault((r.device_tag, r.point_tag), []).append(r)
    ambiguous_points = {k for k, v in by_point.items() if len(v) > 1}
    for r in kept:
        key = (r.device_tag, r.point_tag)
        if key in ambiguous_points:
            others = ", ".join(f"{o.tag}.{o.slot}" for o in by_point[key])
            skipped.append(_Skip(r.row, r.tag, r.slot, "ambiguous",
                                 f"point {key[0]} / {key[1]} is bound to {others}"))
            continue
        if r.device_tag is not None:
            owner = await svc.binding_owner(site.tenant_id, r.device_tag, r.point_tag)
            if owner is not None:
                skipped.append(_Skip(r.row, r.tag, r.slot, "conflict",
                                     f"point {r.device_tag} / {r.point_tag} is already bound "
                                     f"to {owner[1]}.{owner[2]}"))
                continue
        planned[r.tag].slots.append(
            {"slot": r.slot, "device_tag": r.device_tag, "point_tag": r.point_tag}
        )

    systems_used = list(dict.fromkeys(e.system for e in planned.values()))
    kinds = {e.system: e.kind for e in planned.values()}
    skipped.sort(key=lambda s: (s.row, s.slot or ""))
    return {
        "sheet": SHEET,
        "rows_read": len(raw),
        "ignored_columns": ignored,
        "systems": [
            {"name": n, "kind": kinds[n], "reused": n in existing_systems,
             "system_id": existing_systems[n].system_id if n in existing_systems else None}
            for n in systems_used
        ],
        "equipment": [
            {"tag": e.tag, "system": e.system, "equipment_class": e.equipment_class,
             "name": e.name, "design": e.design, "design_units": vocab.units_of(e.design),
             "slots": e.slots, "rows": e.rows, "equipment_id": None}
            for e in planned.values()
        ],
        "skipped": [s.as_dict() for s in skipped],
        "counts": {
            "systems_created": sum(1 for n in systems_used if n not in existing_systems),
            "equipment_created": len(planned),
            "slots_created": sum(len(e.slots) for e in planned.values()),
            "rows_skipped": len({s.row for s in skipped}),
        },
    }


async def apply(svc, site, report: dict[str, Any], *, actor) -> dict[str, Any]:
    """Write `report` — every system, equipment and slot it lists — in ONE commit."""
    from .service import SOURCE_SCHEDULE_IMPORT, _actor_id

    who = _actor_id(actor)
    new_systems: list[SiteSystem] = []
    system_ids: dict[str, str] = {}
    for s in report["systems"]:
        if s["reused"]:
            system_ids[s["name"]] = s["system_id"]
            continue
        row = SiteSystem(tenant_id=site.tenant_id, site_id=site.site_id, name=s["name"],
                         kind=s["kind"], created_by=who, updated_by=who)
        svc.db.add(row)
        new_systems.append(row)
    await svc.db.flush()
    for row in new_systems:
        system_ids[row.name] = row.system_id
    for s in report["systems"]:
        s["system_id"] = system_ids[s["name"]]

    made: list[SiteEquipment] = []
    for e in report["equipment"]:
        row = SiteEquipment(tenant_id=site.tenant_id, site_id=site.site_id,
                            system_id=system_ids[e["system"]], tag=e["tag"], name=e["name"],
                            equipment_class=e["equipment_class"], design=e["design"],
                            created_by=who, updated_by=who)
        svc.db.add(row)
        made.append(row)
    await svc.db.flush()
    for row, e in zip(made, report["equipment"]):
        e["equipment_id"] = row.equipment_id
        for s in e["slots"]:
            svc.db.add(EquipmentPointSlot(tenant_id=site.tenant_id, site_id=site.site_id,
                                          equipment_id=row.equipment_id, slot=s["slot"],
                                          device_tag=s["device_tag"], point_tag=s["point_tag"],
                                          created_by=who, updated_by=who))
    await svc._commit()

    # Published only after the one commit, so a subscriber never hears of a row
    # a rollback took back.
    for row in new_systems:
        await emit(site.tenant_id, "site_system", "created",
                   {"site_id": row.site_id, "system_id": row.system_id, "name": row.name,
                    "kind": row.kind, "description": row.description})
    systems = {s.system_id: s for s in new_systems}
    for row in made:
        system = systems.get(row.system_id)
        snap = await svc.equipment_snapshot(row, system)
        await emit(site.tenant_id, "equipment", "created",
                   {**snap, "source": SOURCE_SCHEDULE_IMPORT})
    # One operator action, one audit entry naming what it made.
    await audit_record(
        svc.db, actor=actor, action="infrastructure.imported", target_type="site",
        target_id=site.site_id,
        meta={"counts": report["counts"], "equipment_tags": [e["tag"] for e in report["equipment"]]},
    )
    return report
